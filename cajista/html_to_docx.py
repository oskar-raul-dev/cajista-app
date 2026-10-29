from __future__ import annotations

import re
from pathlib import Path
from urllib.parse import urlparse

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .css_tokens import StyleTokens
from .docx_ooxml import (
    add_hyperlink,
    disable_spellcheck,
    repeat_table_header,
    set_cell_borders,
    set_cell_shading,
    set_paragraph_borders,
    set_paragraph_shading,
)
from .docx_styles import configure_document

COLOR_RE = re.compile(r"color\s*:\s*#([0-9a-fA-F]{6})")


# Applies font family, size, color and weight to a run, for every script type.
def _set_run_font(run, font_name: str, size: float, color: str, bold=None, italic=None):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), font_name)


# Recursively adds runs for inline HTML (bold, italic, code, links, spans with color).
def _render_inline(paragraph, node, tokens: StyleTokens, bold=False, italic=False, color=None):
    body_font = tokens.font("--font-body", "Arial")
    mono_font = tokens.font("--font-mono", "Consolas")
    text_color = color or tokens.color("--color-text", "1F2937")
    body_size = tokens.points("--font-size-body", 11)

    if isinstance(node, NavigableString):
        if str(node):
            run = paragraph.add_run(str(node))
            _set_run_font(run, body_font, body_size, text_color, bold=bold, italic=italic)
        return
    if not isinstance(node, Tag):
        return

    if node.name in {"strong", "b"}:
        for child in node.children:
            _render_inline(paragraph, child, tokens, True, italic, color)
    elif node.name in {"em", "i"}:
        for child in node.children:
            _render_inline(paragraph, child, tokens, bold, True, color)
    elif node.name == "code":
        run = paragraph.add_run(node.get_text())
        run.style = "Code Inline"
        _set_run_font(run, mono_font, body_size - 0.5, text_color, bold=bold, italic=italic)
    elif node.name == "a":
        href = node.get("href", "")
        add_hyperlink(paragraph, node.get_text(), href, tokens.color("--color-link", "2C3E50"))
    elif node.name == "br":
        paragraph.add_run().add_break()
    elif node.name == "img":
        alt = node.get("alt", "imagen")
        src = node.get("src", "")
        run = paragraph.add_run(f"[{alt}: {src}]")
        run.italic = True
    else:
        style_color = color
        match = COLOR_RE.search(node.get("style", ""))
        if match:
            style_color = match.group(1).upper()
        for child in node.children:
            _render_inline(paragraph, child, tokens, bold, italic, style_color)


# Adds a paragraph built from the tag's inline content.
def _render_paragraph(document, tag: Tag, tokens: StyleTokens, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    for child in tag.children:
        _render_inline(paragraph, child, tokens)
    return paragraph


# Adds a Heading N paragraph with the accent bottom border.
def _render_heading(document, tag: Tag, tokens: StyleTokens):
    level = int(tag.name[1])
    paragraph = document.add_paragraph(style=f"Heading {level}")
    for child in tag.children:
        _render_inline(paragraph, child, tokens)
    # _render_inline sets body size and color on every run; in a heading the
    # size, color and bold of the Heading N style must win.
    for run in paragraph.runs:
        run.font.size = None
        run.font.color.rgb = None
        if not run.bold:
            run.bold = None
    set_paragraph_borders(
        paragraph,
        bottom=(tokens.color("--color-accent", "64748B"), 8, 3),
    )


# Returns the direct <li> children of a list.
def _list_items(tag: Tag):
    return [child for child in tag.children if isinstance(child, Tag) and child.name == "li"]


# Adds list paragraphs, recursing into nested lists up to three levels of style.
def _render_list(document, tag: Tag, tokens: StyleTokens, depth: int = 0):
    ordered = tag.name == "ol"
    style_base = "List Number" if ordered else "List Bullet"
    style = style_base if depth == 0 else f"{style_base} {min(depth + 1, 3)}"
    for item in _list_items(tag):
        paragraph = document.add_paragraph(style=style)
        for child in item.children:
            if isinstance(child, Tag) and child.name in {"ul", "ol"}:
                continue
            _render_inline(paragraph, child, tokens)
        for nested in item.find_all(["ul", "ol"], recursive=False):
            _render_list(document, nested, tokens, depth + 1)


# Adds a Word table with borders, header shading and zebra rows.
def _render_table(document, table_tag: Tag, tokens: StyleTokens, repeat_header: bool):
    rows = table_tag.find_all("tr")
    if not rows:
        return
    col_count = max(len(row.find_all(["th", "td"], recursive=False)) for row in rows)
    table = document.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    border = tokens.color("--color-border", "D1D5DB")
    surface = tokens.color("--color-surface", "F3F4F6")
    surface_alt = tokens.color("--color-surface-alt", "F9FAFB")
    primary = tokens.color("--color-primary", "2C3E50")

    for row_index, source_row in enumerate(rows):
        cells = source_row.find_all(["th", "td"], recursive=False)
        for col_index in range(col_count):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_borders(cell, border)
            if row_index == 0 and any(c.name == "th" for c in cells):
                set_cell_shading(cell, surface)
            elif row_index % 2 == 1:
                set_cell_shading(cell, surface_alt)
            if col_index >= len(cells):
                continue
            source_cell = cells[col_index]
            paragraph = cell.paragraphs[0]
            for child in source_cell.children:
                _render_inline(paragraph, child, tokens, bold=(source_cell.name == "th"), color=(primary if source_cell.name == "th" else None))
    if repeat_header and rows[0].find("th"):
        repeat_table_header(table.rows[0])


# Returns the inline `color:` of a tag, or the inherited color.
def _style_color(tag: Tag, inherited: str) -> str:
    match = COLOR_RE.search(tag.get("style", ""))
    return match.group(1).upper() if match else inherited


# Flattens highlight's spans into (text, color) segments.
def _collect_segments(node, inherited_color: str):
    result: list[tuple[str, str]] = []
    if isinstance(node, NavigableString):
        if str(node):
            result.append((str(node), inherited_color))
        return result
    if not isinstance(node, Tag):
        return result
    current = _style_color(node, inherited_color)
    for child in node.children:
        result.extend(_collect_segments(child, current))
    return result


# Splits (text, color) segments into one list per code line.
def _split_segments_by_line(segments: list[tuple[str, str]]):
    lines: list[list[tuple[str, str]]] = [[]]
    for text, color in segments:
        pieces = text.split("\n")
        for index, piece in enumerate(pieces):
            if piece:
                lines[-1].append((piece, color))
            if index < len(pieces) - 1:
                lines.append([])
    if len(lines) > 1 and not lines[-1]:
        lines.pop()
    return lines or [[]]


# Adds a code block as one shaded paragraph per line, bordered to look like one block.
def _render_code(document, container: Tag, tokens: StyleTokens, disable_proofing: bool):
    pre = container.find("pre") if container.name != "pre" else container
    if pre is None:
        return
    default_color = tokens.color("--color-text", "1F2937")
    segments = _collect_segments(pre, default_color)
    lines = _split_segments_by_line(segments)
    mono_font = tokens.font("--font-mono", "Consolas")
    code_size = tokens.points("--font-size-code", 9.2)
    accent = tokens.color("--color-accent", "64748B")
    border = tokens.color("--color-border", "D1D5DB")
    surface = tokens.color("--color-surface", "F3F4F6")
    is_verbatim = "verbatim-block" in (pre.get("class") or [])
    style_name = "Verbatim Block" if is_verbatim else "Code Block"

    for index, line in enumerate(lines):
        paragraph = document.add_paragraph(style=style_name)
        set_paragraph_shading(paragraph, surface)
        set_paragraph_borders(
            paragraph,
            left=(accent, 18, 6),
            top=(border, 4, 3) if index == 0 else None,
            right=(border, 4, 3),
            bottom=(border, 4, 3) if index == len(lines) - 1 else None,
        )
        if not line:
            paragraph.add_run("")
        for text, color in line:
            run = paragraph.add_run(text)
            _set_run_font(run, mono_font, code_size, color)
            if disable_proofing:
                disable_spellcheck(run)


# Adds each quote paragraph with the Quote style, shading and borders.
def _render_blockquote(document, tag: Tag, tokens: StyleTokens):
    paragraphs = tag.find_all("p", recursive=False) or [tag]
    for paragraph_tag in paragraphs:
        paragraph = _render_paragraph(document, paragraph_tag, tokens, style="Quote")
        set_paragraph_shading(paragraph, tokens.color("--color-surface-alt", "F9FAFB"))
        set_paragraph_borders(
            paragraph,
            left=(tokens.color("--color-accent", "64748B"), 18, 6),
            top=(tokens.color("--color-border", "D1D5DB"), 4, 3),
            right=(tokens.color("--color-border", "D1D5DB"), 4, 3),
            bottom=(tokens.color("--color-border", "D1D5DB"), 4, 3),
        )


# Builds the DOCX by walking the top-level children of `main.document`;
# tags without a handler are skipped.
def html_to_docx(
    html_text: str,
    output_path: Path,
    tokens: StyleTokens,
    word_config: dict,
    reference_docx: Path | None = None,
) -> None:
    document = Document(str(reference_docx)) if reference_docx else Document()
    configure_document(document, tokens, bool(word_config.get("keep_headings_with_next", True)))
    soup = BeautifulSoup(html_text, "lxml")
    root = soup.select_one("main.document") or soup.body
    if root is None:
        raise ValueError("El HTML no contiene body ni main.document.")

    for node in root.children:
        if not isinstance(node, Tag):
            continue
        if node.name in {f"h{i}" for i in range(1, 7)}:
            _render_heading(document, node, tokens)
        elif node.name == "p":
            _render_paragraph(document, node, tokens)
        elif node.name in {"ul", "ol"}:
            _render_list(document, node, tokens)
        elif node.name == "table":
            _render_table(document, node, tokens, bool(word_config.get("repeat_table_header", True)))
        elif node.name == "blockquote":
            _render_blockquote(document, node, tokens)
        elif node.name == "figure" and "code-container" in (node.get("class") or []):
            _render_code(document, node, tokens, bool(word_config.get("disable_spellcheck_in_code", True)))
        elif node.name == "pre":
            _render_code(document, node, tokens, bool(word_config.get("disable_spellcheck_in_code", True)))
        elif node.name == "hr":
            paragraph = document.add_paragraph()
            set_paragraph_borders(paragraph, bottom=(tokens.color("--color-border", "D1D5DB"), 6, 1))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
